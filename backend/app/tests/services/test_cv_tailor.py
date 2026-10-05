import io

from docx import Document

from app.schemas.cv import CVEducation, CVExperience, TailoredCV
from app.services.cv_render import latin1, render_docx, render_pdf, safe_filename
from app.services.cv_tailor import (
    cv_to_text,
    filter_skills,
    keyword_coverage,
    merge_keywords,
    enforce_provenance,
)

ORIGINAL = """Akshay Mohan
Skills: Python, SQL, FastAPI, PostgreSQL, Docker, Git
Software Engineer at Acme Corp, 2022-2024. Built REST APIs with FastAPI.
B.Tech Computer Science, Kerala University"""


def sample_cv(**overrides) -> TailoredCV:
    data = dict(
        full_name="Akshay Mohan",
        headline="Backend Engineer",
        email="a@example.com",
        summary="Backend engineer building APIs with Python and FastAPI.",
        skills=["Python", "FastAPI", "Docker"],
        experience=[CVExperience(role="Software Engineer", company="Acme Corp", start="2022", end="2024",
                                 bullets=["Built REST APIs with FastAPI — serving 3 teams"])],
        education=[CVEducation(degree="B.Tech Computer Science", institution="Kerala University")],
    )
    data.update(overrides)
    return TailoredCV(**data)


def test_filter_skills_drops_skills_not_in_original():
    kept, removed = filter_skills(["Python", "Kubernetes", "Postgres", "Rust"], ORIGINAL)
    assert kept == ["Python", "Postgres"]
    assert removed == ["Kubernetes", "Rust"]


def test_filter_skills_dedupes_case_insensitively():
    kept, _ = filter_skills(["Python", "python", "PYTHON"], ORIGINAL)
    assert kept == ["Python"]


def test_merge_keywords_puts_taxonomy_skills_first_and_dedupes():
    jd = "We need Python, Kubernetes and strong stakeholder management."
    merged = merge_keywords(["stakeholder management", "python"], jd)
    assert merged[:2] == ["Kubernetes", "Python"]
    assert merged.count("Python") + merged.count("python") == 1
    assert "stakeholder management" in merged


def test_keyword_coverage_counts_aliases():
    matched, coverage = keyword_coverage(["PostgreSQL", "Kubernetes"], "Worked with Postgres daily")
    assert matched == ["PostgreSQL"]
    assert coverage == 0.5


def test_provenance_blanks_unknown_employer_and_institution():
    cv = sample_cv(
        experience=[CVExperience(role="Engineer", company="(Current Employer)")],
        education=[CVEducation(degree="B.Tech", institution="IIT Bombay")],
    )
    warnings = enforce_provenance(cv, ORIGINAL)
    assert cv.experience[0].company is None and cv.education[0].institution is None
    assert any("(Current Employer)" in w for w in warnings)
    assert any("IIT Bombay" in w for w in warnings)


def test_provenance_keeps_verified_facts_silently():
    cv = sample_cv()
    assert enforce_provenance(cv, ORIGINAL) == []
    assert cv.experience[0].company == "Acme Corp"


def test_cv_to_text_includes_bullets_and_skills():
    text = cv_to_text(sample_cv())
    assert "Built REST APIs" in text and "Docker" in text


def test_latin1_replaces_smart_punctuation():
    assert latin1("A — “quoted” • item") == 'A - "quoted" - item'


def test_render_pdf_produces_pdf_bytes():
    data = render_pdf(sample_cv())
    assert data.startswith(b"%PDF")


def test_render_docx_contains_cv_text():
    doc = Document(io.BytesIO(render_docx(sample_cv())))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Akshay Mohan" in text
    assert "Built REST APIs with FastAPI" in text


def test_safe_filename_strips_unsafe_characters():
    assert safe_filename(sample_cv(full_name="Akshay / Mohan!"), "pdf") == "Akshay_Mohan_CV.pdf"


def test_umbrella_skill_allowed_when_child_skill_present():
    kept, removed = filter_skills(["Machine Learning"], "Trained PyTorch models")
    assert kept == ["Machine Learning"] and removed == []


def test_umbrella_skill_not_counted_as_literal_ats_match():
    matched, _ = keyword_coverage(["Machine Learning"], "Trained PyTorch models")
    assert matched == []


def test_warns_when_new_technology_appears_in_a_bullet():
    cv = sample_cv(experience=[CVExperience(role="Software Engineer", company="Acme Corp",
                                            bullets=["Deployed services on Kubernetes"])])
    warnings = enforce_provenance(cv, ORIGINAL)
    assert any("Kubernetes" in w for w in warnings)


def test_extract_json_handles_surrounding_text():
    from app.services.cv_tailor import extract_json
    assert extract_json('Here you go:\n{"cv": {"full_name": "A"}}\nDone') == {"cv": {"full_name": "A"}}
    assert extract_json("no json here") is None
    assert extract_json("") is None


def test_cv_accepts_null_company_and_numeric_year():
    cv = TailoredCV.model_validate({
        "full_name": "A", "links": None,
        "experience": [{"role": "Engineer", "company": None, "bullets": ["Built APIs"]}],
        "education": [{"degree": "B.Tech", "institution": None, "year": 2022}],
    })
    assert cv.links == [] and cv.experience[0].company is None and cv.education[0].year == "2022"
    assert render_pdf(cv).startswith(b"%PDF")
    text = "\n".join(p.text for p in Document(io.BytesIO(render_docx(cv))).paragraphs)
    assert "None" not in text and "Engineer" in text


class _FakeGroq:
    """Replays a scripted sequence of responses/exceptions for chat.completions.create."""

    def __init__(self, script):
        self.calls = []
        script = list(script)

        def create(**kwargs):
            self.calls.append(kwargs)
            item = script.pop(0)
            if isinstance(item, Exception):
                raise item
            message = type("M", (), {"content": item})()
            return type("R", (), {"choices": [type("C", (), {"message": message})()]})()

        self.chat = type("Chat", (), {"completions": type("Comp", (), {"create": staticmethod(create)})()})()


def _patch_groq(monkeypatch, fake):
    import app.services.cv_tailor as mod
    monkeypatch.setattr(mod, "GROQ_API_KEY", "test-key")
    monkeypatch.setattr(mod, "Groq", lambda api_key: fake)
    return mod


def test_rate_limit_is_not_retried_and_reports_busy(monkeypatch):
    import httpx
    import pytest
    from groq import RateLimitError
    request = httpx.Request("POST", "https://api.groq.com")
    fake = _FakeGroq([RateLimitError("limit", response=httpx.Response(429, request=request), body=None)])
    mod = _patch_groq(monkeypatch, fake)
    with pytest.raises(mod.CVTailorBusyError):
        mod.call_llm(ORIGINAL, "Engineer", "Acme", "x" * 100)
    assert len(fake.calls) == 1


def test_json_mode_failure_falls_back_to_plain_mode(monkeypatch):
    import httpx
    from groq import BadRequestError
    request = httpx.Request("POST", "https://api.groq.com")
    bad = BadRequestError("json_validate_failed", response=httpx.Response(400, request=request), body=None)
    fake = _FakeGroq([bad, bad, 'Sure: {"cv": {"full_name": "A"}}'])
    mod = _patch_groq(monkeypatch, fake)
    assert mod.call_llm(ORIGINAL, "Engineer", "Acme", "x" * 100) == {"cv": {"full_name": "A"}}
    assert "response_format" not in fake.calls[-1]
