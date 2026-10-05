import io
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF

from app.schemas.cv import TailoredCV

_PUNCTUATION = {
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "-", "•": "-", "…": "...",
    " ": " ", " ": " ", "​": "",
}


def latin1(text: str | None) -> str:
    text = "".join(_PUNCTUATION.get(ch, ch) for ch in (text or ""))
    text = unicodedata.normalize("NFKD", text)
    return text.encode("latin-1", "ignore").decode("latin-1")


def contact_line(cv: TailoredCV) -> str:
    return " | ".join(p for p in [cv.email, cv.phone, cv.location, *cv.links] if p)


def joined(*parts: str | None, sep: str = " - ") -> str:
    return sep.join(p for p in parts if p)


def date_range(start: str | None, end: str | None) -> str:
    return joined(start, end)


def safe_filename(cv: TailoredCV, extension: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9]+", "_", cv.full_name or "CV").strip("_") or "CV"
    return f"{stem}_CV.{extension}"


class _PDF(FPDF):
    def heading(self, title: str) -> None:
        self.ln(3)
        self.set_font("Helvetica", "B", 11)
        self.cell(0, 6, title.upper(), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(150, 150, 150)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(1.5)

    def para(self, value: str, style: str = "", size: int = 10) -> None:
        self.set_font("Helvetica", style, size)
        self.multi_cell(0, 5, latin1(value), new_x="LMARGIN", new_y="NEXT")

    def bullet(self, value: str) -> None:
        self.set_font("Helvetica", "", 10)
        self.set_x(self.l_margin + 3)
        self.multi_cell(0, 5, "- " + latin1(value), new_x="LMARGIN", new_y="NEXT")


def render_pdf(cv: TailoredCV) -> bytes:
    pdf = _PDF(format="A4")
    pdf.set_margins(16, 14, 16)
    pdf.set_auto_page_break(True, margin=14)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 9, latin1(cv.full_name or "Curriculum Vitae"), new_x="LMARGIN", new_y="NEXT")
    if cv.headline:
        pdf.para(cv.headline, size=11)
    if contact_line(cv):
        pdf.para(contact_line(cv), size=9)

    if cv.summary:
        pdf.heading("Summary")
        pdf.para(cv.summary)
    if cv.skills:
        pdf.heading("Skills")
        pdf.para(", ".join(cv.skills))
    if cv.experience:
        pdf.heading("Experience")
        for job in cv.experience:
            pdf.para(joined(job.role, job.company), style="B")
            meta = " | ".join(p for p in [date_range(job.start, job.end), job.location] if p)
            if meta:
                pdf.para(meta, size=9)
            for line in job.bullets:
                pdf.bullet(line)
            pdf.ln(1.5)
    if cv.projects:
        pdf.heading("Projects")
        for project in cv.projects:
            if project.name:
                pdf.para(project.name, style="B")
            for line in project.bullets:
                pdf.bullet(line)
            pdf.ln(1.5)
    if cv.education:
        pdf.heading("Education")
        for edu in cv.education:
            pdf.para(joined(edu.degree, edu.institution), style="B")
            meta = " | ".join(p for p in [edu.year, edu.details] if p)
            if meta:
                pdf.para(meta, size=9)
    if cv.certifications:
        pdf.heading("Certifications")
        for cert in cv.certifications:
            pdf.bullet(cert)

    return bytes(pdf.output())


def render_docx(cv: TailoredCV) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(cv.full_name or "Curriculum Vitae")
    run.bold = True
    run.font.size = Pt(18)
    if cv.headline:
        doc.add_paragraph(cv.headline)
    if contact_line(cv):
        doc.add_paragraph(contact_line(cv))

    def section(name: str) -> None:
        doc.add_heading(name, level=2)

    def bold_line(value: str) -> None:
        doc.add_paragraph().add_run(value).bold = True

    if cv.summary:
        section("Summary")
        doc.add_paragraph(cv.summary)
    if cv.skills:
        section("Skills")
        doc.add_paragraph(", ".join(cv.skills))
    if cv.experience:
        section("Experience")
        for job in cv.experience:
            bold_line(joined(job.role, job.company))
            meta = " | ".join(p for p in [date_range(job.start, job.end), job.location] if p)
            if meta:
                doc.add_paragraph(meta)
            for line in job.bullets:
                doc.add_paragraph(line, style="List Bullet")
    if cv.projects:
        section("Projects")
        for project in cv.projects:
            if project.name:
                bold_line(project.name)
            for line in project.bullets:
                doc.add_paragraph(line, style="List Bullet")
    if cv.education:
        section("Education")
        for edu in cv.education:
            bold_line(joined(edu.degree, edu.institution))
            meta = " | ".join(p for p in [edu.year, edu.details] if p)
            if meta:
                doc.add_paragraph(meta)
    if cv.certifications:
        section("Certifications")
        for cert in cv.certifications:
            doc.add_paragraph(cert, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
